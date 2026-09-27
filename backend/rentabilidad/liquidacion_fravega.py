"""Lector de la liquidación quincenal de Frávega (Seller Center → Excel).

Es la única fuente real de comisión y fee logístico de Frávega: la API de
Ecom trae la venta pero no lo que cobra Frávega (el pago viene sin detalle).
Decisión de Maxx (2026-09-27): mientras no llegó la liquidación de la
quincena, las órdenes de Frávega se calculan con la comisión base estimada;
cuando llega, se reemplazan por estos valores reales y se recalcula la
rentabilidad (`persistencia.aplicar_liquidacion_fravega`).

Estructura real del archivo (verificada contra `Liq-al-15062026.xlsx`,
período 01/06→15/06/2026, y `liquidacion actual.xlsx`, 16/07→31/07/2026):

- Pestaña `Totales`: pares etiqueta/valor (`Liquidación desde:`,
  `Total Comisiones  (1)`, `Total Servicios logísticos - Fee (2)`...) —
  los totales vienen en negativo (son descuentos al vendedor).
- Pestaña `Detalle de Operaciones`: una fila por unidad vendida. La primera
  fila de cada orden trae `Orden`; las unidades siguientes de la misma
  orden vienen con `Orden` vacío (se arrastra la última orden vista). El
  `Fee logistico` va una sola vez por orden.
- `Tipo de operacion` = `Facturación` o `Devolución`. La devolución trae los
  mismos importes en positivo y revierte todo (valor, comisión y fee):
  verificado contra `Totales` — Σ Facturación − Σ Devolución da exacto
  "Total Facturado neto de cancelaciones", "Total Comisiones" y "Total
  Servicios logísticos - Fee".

Antes de devolver nada se controla que la suma del detalle reconstruya los
totales de la pestaña `Totales`: si no cuadra, el archivo no se usa
(`LiquidacionFravegaInvalida`) — aplicar una comisión mal leída a la
rentabilidad es peor que dejarla estimada.
"""
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from typing import Callable

_TAB_TOTALES = "Totales"
_TAB_DETALLE = "Detalle de Operaciones"

_COL_ORDEN = "Orden"
_COL_TIPO = "Tipo de operacion"
_COL_FEE = "Fee logistico"
_COL_VALOR_SKU = "Valor del sku"
_COL_COMISION = "Valor total Comisiones"

_TIPO_FACTURACION = "Facturación"
_TIPO_DEVOLUCION = "Devolución"

# Etiquetas de `Totales`, comparadas sin espacios repetidos ni al final
# (el archivo real trae `Total Comisiones  (1)`, `CUIT: `...).
_TOTAL_DESDE = "Liquidación desde:"
_TOTAL_HASTA = "Liquidación hasta:"
_TOTAL_FACTURADO = "Total Facturado c/ IVA neto de cancelaciones"
_TOTAL_COMISIONES = "Total Comisiones (1)"
_TOTAL_FEE = "Total Servicios logísticos - Fee (2)"

_TOLERANCIA = Decimal("0.05")


class LiquidacionFravegaInvalida(ValueError):
    """El archivo no tiene la forma esperada o su detalle no reconstruye
    los totales de la propia liquidación."""


@dataclass
class OrdenLiquidada:
    orden: str  # ej. "v90781066frvg-01"
    valor_sku: Decimal = Decimal(0)
    comision: Decimal = Decimal(0)
    fee_logistico: Decimal = Decimal(0)

    @property
    def cancelada(self) -> bool:
        """Revertida entera por una devolución: no queda venta neta."""
        return self.valor_sku <= 0


@dataclass
class LiquidacionFravega:
    desde: date | None
    hasta: date | None
    ordenes: dict[str, OrdenLiquidada] = field(default_factory=dict)


def _decimal(v) -> Decimal:
    if v in (None, ""):
        return Decimal(0)
    return Decimal(str(v))


def _etiqueta(v) -> str:
    return " ".join(str(v).split()) if v is not None else ""


def _fecha(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if not v:
        return None
    return datetime.strptime(str(v).strip(), "%d/%m/%Y").date()


def leer_hojas_excel(path: str) -> tuple[dict[str, object], list[dict]]:
    """Import perezoso de openpyxl (mismo criterio que `ingesta_ecom`).
    Devuelve (totales etiqueta→valor, filas del detalle como dicts)."""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if _TAB_TOTALES not in wb.sheetnames or _TAB_DETALLE not in wb.sheetnames:
        raise LiquidacionFravegaInvalida(
            f"Faltan las pestañas '{_TAB_TOTALES}' / '{_TAB_DETALLE}' — ¿es una liquidación de Frávega?"
        )
    totales = {
        _etiqueta(fila[0]): fila[1]
        for fila in wb[_TAB_TOTALES].iter_rows(values_only=True)
        if fila and len(fila) > 1 and fila[0]
    }
    filas = list(wb[_TAB_DETALLE].iter_rows(values_only=True))
    encabezados = [_etiqueta(h) for h in filas[0]] if filas else []
    detalle = [dict(zip(encabezados, fila)) for fila in filas[1:]]
    return totales, detalle


LeerHojas = Callable[[str], tuple[dict[str, object], list[dict]]]


def procesar_liquidacion(totales: dict[str, object], detalle: list[dict]) -> LiquidacionFravega:
    """Agrupa el detalle por orden, neto de devoluciones, y lo valida contra
    `Totales`."""
    liquidacion = LiquidacionFravega(desde=_fecha(totales.get(_TOTAL_DESDE)), hasta=_fecha(totales.get(_TOTAL_HASTA)))
    orden_actual = None
    for fila in detalle:
        tipo = _etiqueta(fila.get(_COL_TIPO))
        if not tipo:
            continue  # filas vacías de relleno al final de la hoja
        if fila.get(_COL_ORDEN):
            orden_actual = _etiqueta(fila[_COL_ORDEN])
        if orden_actual is None:
            raise LiquidacionFravegaInvalida("Hay una fila de detalle antes de la primera orden.")
        if tipo == _TIPO_FACTURACION:
            signo = Decimal(1)
        elif tipo == _TIPO_DEVOLUCION:
            signo = Decimal(-1)
        else:
            raise LiquidacionFravegaInvalida(f"Tipo de operación no relevado: {tipo!r} (orden {orden_actual}).")

        orden = liquidacion.ordenes.setdefault(orden_actual, OrdenLiquidada(orden=orden_actual))
        orden.valor_sku += signo * _decimal(fila.get(_COL_VALOR_SKU))
        orden.comision += signo * _decimal(fila.get(_COL_COMISION))
        orden.fee_logistico += signo * _decimal(fila.get(_COL_FEE))

    _validar_contra_totales(liquidacion, totales)
    return liquidacion


def _validar_contra_totales(liquidacion: LiquidacionFravega, totales: dict[str, object]) -> None:
    controles = (
        (_TOTAL_FACTURADO, sum((o.valor_sku for o in liquidacion.ordenes.values()), Decimal(0)), Decimal(1)),
        (_TOTAL_COMISIONES, sum((o.comision for o in liquidacion.ordenes.values()), Decimal(0)), Decimal(-1)),
        (_TOTAL_FEE, sum((o.fee_logistico for o in liquidacion.ordenes.values()), Decimal(0)), Decimal(-1)),
    )
    for etiqueta, calculado, signo in controles:
        if etiqueta not in totales:
            raise LiquidacionFravegaInvalida(f"Falta '{etiqueta}' en la pestaña {_TAB_TOTALES}.")
        informado = signo * _decimal(totales[etiqueta])
        if abs(calculado - informado) > _TOLERANCIA:
            raise LiquidacionFravegaInvalida(
                f"El detalle no reconstruye '{etiqueta}': suma {calculado} vs informado {informado}."
            )


class LiquidacionFravegaAdapter:
    def __init__(self, leer_hojas: LeerHojas | None = None):
        self._leer_hojas = leer_hojas or leer_hojas_excel

    def procesar(self, path: str) -> LiquidacionFravega:
        totales, detalle = self._leer_hojas(path)
        return procesar_liquidacion(totales, detalle)
