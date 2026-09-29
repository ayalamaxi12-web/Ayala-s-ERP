"""Frávega: comisión estimada hasta la liquidación quincenal, real después
(decisión de Maxx, 2026-09-27). Fixtures con la forma real de la liquidación
`Liq-al-15062026.xlsx` (01/06→15/06/2026) y órdenes reales del 01/06/2026."""
from datetime import date
from decimal import Decimal

import pytest

from rentabilidad.adapters import (
    ClasificacionProvider,
    IvaProvider,
    MargenObjetivoProvider,
    StockProvider,
    VinculacionProvider,
)
from rentabilidad.ingesta_ecom import FilaEcom, ResultadoIngestaEcom
from rentabilidad.liquidacion_fravega import (
    LiquidacionFravegaAdapter,
    LiquidacionFravegaInvalida,
    procesar_liquidacion,
)
from rentabilidad.models import LiquidacionFravega, VentaEcom
from rentabilidad.persistencia import (
    OBSERVACION_CANCELADA_EN_FRAVEGA,
    aplicar_liquidacion_fravega,
    guardar_cierre_ecom,
    ventas_fravega_estimadas,
)

TC = Decimal("1540")


def _cerca(a, b, tol=Decimal("0.01")):
    return abs(Decimal(a) - Decimal(b)) <= tol


def _totales(desde="1/6/2026", hasta="15/6/2026", facturado=0, comisiones=0, fee=0):
    """Pestaña `Totales` real: los descuentos al vendedor vienen en negativo
    y las etiquetas traen espacios de más (se normalizan al leer)."""
    return {
        "Liquidación desde:": desde, "Liquidación hasta:": hasta,
        "Total Facturado c/ IVA neto de cancelaciones": facturado,
        "Total Comisiones (1)": -comisiones,
        "Total Servicios logísticos - Fee (2)": -fee,
    }


def _op(orden, tipo, valor, comision, fee):
    return {"Orden": orden, "Tipo de operacion": tipo, "Valor del sku": valor,
            "Valor total Comisiones": comision, "Fee logistico": fee}


# Detalle real (recortado) — 1387346 = v90781066frvg-01 (comisión base 15%
# + financiera 16%), 1387291 = v90777766frvg-01 (facturada el 02/06 y
# devuelta el 05/06 dentro de la misma liquidación).
_DETALLE = [
    _op("v90781066frvg-01", "Facturación", 760537.19, 235766.5289, 14639),
    _op("v90777766frvg-01", "Facturación", 28899, 4334.85, 1016),
    _op("v90777766frvg-01", "Devolución", 28899, 4334.85, 1016),
    # dos unidades de la misma orden: la segunda fila viene sin "Orden"
    _op("v90711591frvg-01", "Facturación", 8479, 1271.85, 2755),
    _op(None, "Facturación", 8479, 1271.85, None),
    {"Orden": None, "Tipo de operacion": None},  # relleno al final de la hoja
]


def _liquidacion():
    return procesar_liquidacion(
        _totales(facturado=760537.19 + 2 * 8479, comisiones=235766.5289 + 2 * 1271.85, fee=14639 + 2755),
        _DETALLE,
    )


# ── Lector de la liquidación ──

def test_liquidacion_agrupa_por_orden_arrastrando_la_orden_en_filas_sin_numero():
    liq = _liquidacion()
    assert liq.desde == date(2026, 6, 1) and liq.hasta == date(2026, 6, 15)
    orden = liq.ordenes["v90711591frvg-01"]
    assert orden.valor_sku == Decimal("16958")
    assert orden.comision == Decimal("2543.70")
    assert orden.fee_logistico == Decimal("2755")  # el fee va una sola vez por orden


def test_liquidacion_la_devolucion_revierte_valor_comision_y_fee():
    orden = _liquidacion().ordenes["v90777766frvg-01"]
    assert orden.valor_sku == 0 and orden.comision == 0 and orden.fee_logistico == 0
    assert orden.cancelada


def test_liquidacion_que_no_reconstruye_sus_totales_no_se_usa():
    with pytest.raises(LiquidacionFravegaInvalida, match="Total Comisiones"):
        procesar_liquidacion(_totales(facturado=760537.19, comisiones=1, fee=14639), _DETALLE[:1])


def test_liquidacion_con_tipo_de_operacion_desconocido_no_se_usa():
    with pytest.raises(LiquidacionFravegaInvalida, match="no relevado"):
        procesar_liquidacion(_totales(), [_op("v1frvg-01", "Ajuste raro", 1, 0, 0)])


def test_adapter_lee_con_la_funcion_inyectada():
    liq = LiquidacionFravegaAdapter(leer_hojas=lambda path: (
        _totales(facturado=760537.19, comisiones=235766.5289, fee=14639), _DETALLE[:1],
    )).procesar("x.xlsx")
    assert list(liq.ordenes) == ["v90781066frvg-01"]


# ── Persistencia: estimado → real ──

def _sin_clasificar():
    return dict(
        clasificacion_provider=ClasificacionProvider(sheet_id=None),
        vinculacion_provider=VinculacionProvider(sheet_id=None),
        stock_provider=StockProvider(sheet_id=None),
        margen_provider=MargenObjetivoProvider(sheet_ids={}, sheet_master_id=None),
    )


def _fila_fravega(numero_orden="1387346", orden_externa="v90781066frvg-01", **overrides):
    base = dict(
        numero_orden=numero_orden, skus_vendidos="PLANCHA-SUB-30X38-5EN1", canal_de_venta="Fravega",
        estado_pago="Cobrado", costo_sin_iva=Decimal("96.34"), comision_venta=Decimal(0),
        costo_envio=Decimal(0), precio_sin_iva=Decimal("688268.95"), precio_final=Decimal("760537.19"),
        tc=TC, orden_externa=orden_externa, origen_comision="ESTIMADO_FRAVEGA",
    )
    base.update(overrides)
    return FilaEcom(**base)


def _guardar(db, *filas, periodo="2026-06-01_2026-06-01"):
    ingesta = ResultadoIngestaEcom(lineas=list(filas), excluidas_por_estado_pago=[], incidencias_costo=[])
    guardar_cierre_ecom(db, periodo, ingesta, IvaProvider(), **_sin_clasificar())
    db.flush()


def _venta(db, numero_orden):
    return db.query(VentaEcom).filter_by(numero_orden=numero_orden).one()


# IVA que Frávega factura aparte sobre comisión y fee: se suma al costo
# (Maxx, 2026-09-29 — comisión completa, mismo criterio que ML/MP).
IVA_CARGOS = Decimal("1.21")


def _rentabilidad(precio_sin_iva, precio_final, comision, envio, costo_usd):
    q, v = Decimal(precio_sin_iva), Decimal(precio_final)
    return q - comision - envio - v * Decimal("0.012") - q * Decimal("0.05") - Decimal("149.12") - Decimal(costo_usd) * TC


def test_sin_liquidacion_fravega_se_guarda_con_la_comision_base_estimada(db_session):
    _guardar(db_session, _fila_fravega())
    venta = _venta(db_session, "1387346")
    assert venta.origen_comision == "ESTIMADO_FRAVEGA"
    assert venta.comision_venta == Decimal("760537.19") * Decimal("0.15") * IVA_CARGOS
    assert venta.costo_envio == 0
    assert venta.orden_externa == "v90781066frvg-01"
    assert _cerca(venta.rentabilidad, _rentabilidad("688268.95", "760537.19", Decimal("114080.5785") * IVA_CARGOS, 0, "96.34"))
    assert [v.numero_orden for v in ventas_fravega_estimadas(db_session)] == ["1387346"]


def test_al_cargar_la_liquidacion_se_reemplaza_por_el_real_y_se_recalcula(db_session):
    _guardar(db_session, _fila_fravega())
    resultado = aplicar_liquidacion_fravega(db_session, _liquidacion(), archivo="Liq-al-15062026.xlsx")

    venta = _venta(db_session, "1387346")
    assert venta.origen_comision == "LIQUIDACION_FRAVEGA"
    assert venta.comision_venta == Decimal("235766.5289") * IVA_CARGOS
    assert venta.costo_envio == Decimal("14639") * IVA_CARGOS
    # 245.810,81 — el número del 01/06/2026 validado con Maxx, con comisión y
    # fee sin IVA — menos el 21% de IVA sobre ambos que ahora se descuenta:
    # 245.810,81 − 0,21 × (235.766,53 + 14.639) = 193.225,65
    assert _cerca(venta.rentabilidad, Decimal("193225.65"), tol=Decimal("0.5"))
    assert _cerca(venta.rentabilidad_usd, venta.rentabilidad / TC)
    assert resultado.ventas_actualizadas == ["1387346"]
    assert ventas_fravega_estimadas(db_session) == []


def test_recargar_la_misma_liquidacion_no_duplica_ni_cambia_el_numero(db_session):
    _guardar(db_session, _fila_fravega())
    aplicar_liquidacion_fravega(db_session, _liquidacion())
    antes = _venta(db_session, "1387346").rentabilidad
    aplicar_liquidacion_fravega(db_session, _liquidacion())
    assert _venta(db_session, "1387346").rentabilidad == antes
    assert db_session.query(LiquidacionFravega).filter_by(orden="v90781066frvg-01").count() == 1


def test_un_cierre_guardado_despues_de_la_liquidacion_ya_toma_el_real(db_session):
    aplicar_liquidacion_fravega(db_session, _liquidacion())
    _guardar(db_session, _fila_fravega())  # ej. se re-guarda el día ya liquidado
    venta = _venta(db_session, "1387346")
    assert venta.origen_comision == "LIQUIDACION_FRAVEGA"
    assert venta.comision_venta == Decimal("235766.5289") * IVA_CARGOS


def test_cancelada_en_fravega_pero_cobrada_en_ecom_no_se_fuerza_se_observa(db_session):
    # 1387291 real: la liquidación la revierte entera; Ecom la tiene cobrada.
    fila = _fila_fravega(
        numero_orden="1387291", orden_externa="v90777766frvg-01", skus_vendidos="TVS-ARM-14-55I-400X400",
        costo_sin_iva=Decimal("5.2"), precio_sin_iva=Decimal("23883.47"), precio_final=Decimal("28899"),
    )
    _guardar(db_session, fila)
    estimada = _venta(db_session, "1387291").rentabilidad

    resultado = aplicar_liquidacion_fravega(db_session, _liquidacion())

    venta = _venta(db_session, "1387291")
    assert resultado.canceladas_observadas == ["1387291"]
    assert venta.observacion == OBSERVACION_CANCELADA_EN_FRAVEGA
    assert venta.excluido is False  # no se fuerza la exclusión
    assert venta.origen_comision == "ESTIMADO_FRAVEGA"
    assert venta.rentabilidad == estimada


def test_devolucion_en_una_quincena_posterior_suma_contra_la_venta(db_session):
    _guardar(db_session, _fila_fravega())
    aplicar_liquidacion_fravega(db_session, _liquidacion())
    segunda = procesar_liquidacion(
        _totales(desde="16/6/2026", hasta="30/6/2026", facturado=-760537.19, comisiones=-235766.5289, fee=-14639),
        [_op("v90781066frvg-01", "Devolución", 760537.19, 235766.5289, 14639)],
    )
    resultado = aplicar_liquidacion_fravega(db_session, segunda)
    venta = _venta(db_session, "1387346")
    assert resultado.canceladas_observadas == ["1387346"]
    assert venta.observacion == OBSERVACION_CANCELADA_EN_FRAVEGA
    assert db_session.query(LiquidacionFravega).filter_by(orden="v90781066frvg-01").count() == 2


def test_ordenes_de_la_liquidacion_sin_venta_guardada_se_cuentan_aparte(db_session):
    _guardar(db_session, _fila_fravega())
    resultado = aplicar_liquidacion_fravega(db_session, _liquidacion())
    assert resultado.ordenes_en_liquidacion == 3
    assert resultado.sin_venta_en_ecom == 2


def test_las_ventas_que_no_son_fravega_no_se_tocan(db_session):
    ml = _fila_fravega(numero_orden="1", orden_externa=None, origen_comision="API", canal_de_venta="Mercadolibre",
                       comision_venta=Decimal("100"), costo_envio=Decimal("50"))
    _guardar(db_session, ml)
    venta = _venta(db_session, "1")
    assert venta.comision_venta == Decimal("100") and venta.costo_envio == Decimal("50")
    assert venta.origen_comision == "API"


# ── Endpoint: subir el .xlsx tal cual sale de Seller Center ──

def _xlsx_liquidacion(path):
    import openpyxl

    wb = openpyxl.Workbook()
    tot = wb.active
    tot.title = "Totales"
    for fila in [
        ("Seller:", "Global"), ("Liquidación desde:", "1/6/2026"), ("Liquidación hasta:", "15/6/2026"),
        ("Totales", "Valores $"), ("Total Facturado c/ IVA neto de cancelaciones", 760537.19),
        ("Total Comisiones  (1)", -235766.5289), ("Total Servicios logísticos - Fee (2)", -14639),
    ]:
        tot.append(fila)
    det = wb.create_sheet("Detalle de Operaciones")
    det.append(["Orden", "Id seq", "Tipo de operacion", "Fecha de operacion", "Fecha de compra", "Fee logistico",
                "Envío", "Sku", "Nombre del Sku", "Valor del sku", "Comisión Base", "Valor Comisión Base",
                "Comisión Comercial", "Valor Comisión Comercial", "Comisión Financiera", "Valor Comisión Financiera",
                "Total Comisiones", "Valor total Comisiones", "Valor total Descuento Comercial", "Nro de recibo"])
    det.append(["v90781066frvg-01", "1", "Facturación", "2/6/2026", "1/6/2026", 14639, 0, "1", "Plancha", 760537.19,
                0.15, 114080.5785, 0, 0, 0.16, 121685.9504, 0.31, 235766.5289, 0, None])
    wb.save(path)


@pytest.fixture()
def db_multihilo():
    """Como `db_session`, pero usable desde el hilo del TestClient (SQLite en
    memoria compartido entre hilos)."""
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from sqlalchemy.pool import StaticPool

    from rentabilidad import seed
    from rentabilidad.db import Base

    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool, future=True)
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine, future=True, expire_on_commit=False)()
    seed.seed(session)
    session.commit()
    try:
        yield session
    finally:
        session.close()


def test_endpoint_carga_la_liquidacion_y_actualiza_las_ventas(db_multihilo, monkeypatch, tmp_path):
    db_session = db_multihilo
    from contextlib import contextmanager

    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from rentabilidad import api

    @contextmanager
    def _sesion():
        yield db_session
        db_session.flush()

    monkeypatch.setattr(api, "sesion", _sesion)
    _guardar(db_session, _fila_fravega())
    path = tmp_path / "Liq-al-15062026.xlsx"
    _xlsx_liquidacion(path)

    app = FastAPI()
    app.include_router(api.router)
    with open(path, "rb") as f:
        r = TestClient(app).post("/rentabilidad/fravega/liquidacion", files={"archivo": (path.name, f)})

    assert r.status_code == 200, r.text
    cuerpo = r.json()
    assert cuerpo["desde"] == "2026-06-01" and cuerpo["ventas_actualizadas"] == 1
    assert cuerpo["fravega_pendientes"] == 0
    assert _venta(db_session, "1387346").origen_comision == "LIQUIDACION_FRAVEGA"
    assert TestClient(app).get("/rentabilidad/fravega/pendientes").json() == []


def test_endpoint_rechaza_una_liquidacion_que_no_cuadra(db_session, monkeypatch, tmp_path):
    import openpyxl
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from rentabilidad import api

    path = tmp_path / "liq.xlsx"
    _xlsx_liquidacion(path)
    wb = openpyxl.load_workbook(path)
    wb["Totales"]["B6"] = -1  # Total Comisiones que no coincide con el detalle
    wb.save(path)

    app = FastAPI()
    app.include_router(api.router)
    with open(path, "rb") as f:
        r = TestClient(app).post("/rentabilidad/fravega/liquidacion", files={"archivo": (path.name, f)})
    assert r.status_code == 422
    assert "Total Comisiones" in r.json()["detail"]


# ── OnCity: comisión estimada (canal manual en Ecom, Maxx 2026-09-29) ──

def test_oncity_se_guarda_con_la_comision_estimada(db_session):
    fila = _fila_fravega(numero_orden="1424588", orden_externa=None, canal_de_venta="OnCity",
                         origen_comision="ESTIMADO_ONCITY", precio_final=Decimal("17481"),
                         precio_sin_iva=Decimal("14447.11"), costo_sin_iva=Decimal("2.1"))
    _guardar(db_session, fila)
    venta = _venta(db_session, "1424588")
    assert venta.comision_venta == Decimal("17481") * Decimal("0.15") * IVA_CARGOS
    assert venta.origen_comision == "ESTIMADO_ONCITY"
    assert _cerca(venta.rentabilidad, _rentabilidad("14447.11", "17481", Decimal("2622.15") * IVA_CARGOS, 0, "2.1"))
