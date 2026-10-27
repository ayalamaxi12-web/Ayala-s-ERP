"""Carga de ofertas por Excel («Precios y ofertas ML» → «Cargar un lote por Excel»).

Un Excel por cuenta con 4 columnas: SKU, MLA, precio tachado, precio final. Tres pasos:

1. `leer_excel` — lee y valida el formato (sin tocar Mercado Libre).
2. `iniciar_analisis` — job de SOLO LECTURA: para cada MLA mira la publicación y sus promociones, arma las alertas
   (bloqueos = «no se puede», avisos = «ojo») y la lista de campañas propias con cuántos MLA de los subidos pueden
   participar en cada una.
3. `iniciar_aplicacion` — job que ESCRIBE en Mercado Libre: mete en la campaña elegida las filas que se pueden, con
   el tachado y el final EXACTOS del Excel (`activar_en_campana_tradicional(..., tachado_exacto=...)`), una por una,
   y devuelve el resultado de cada fila. Las que no entran quedan pendientes para probar en otra campaña.

El margen solo avisa (decisión de Maxx): nunca bloquea.
"""
from __future__ import annotations

import io
import re
import threading
import unicodedata
from concurrent.futures import ThreadPoolExecutor
from decimal import Decimal, InvalidOperation
from typing import Callable

from ml_ofertas import (
    MLOfertasClient, MLOfertasEscritura, ParametrosMargen, _cuotas_sin_interes, _sku_de_item,
    activar_en_campana_tradicional, calcular_margen_oferta,
)

MAX_FILAS = 500
DESCUENTO_MIN_PCT = Decimal(10)   # campañas propias: ML exige entre 10 % y 80 % de descuento
DESCUENTO_MAX_PCT = Decimal(80)
_MLA = re.compile(r"MLA-?(\d{6,})", re.I)

_jobs: dict[str, dict] = {}


def estado_job(job_id: str) -> dict | None:
    return _jobs.get(job_id)


def _norm(t) -> str:
    t = unicodedata.normalize("NFKD", str(t or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", t).strip()


def _columna(encabezado: str) -> str | None:
    n = _norm(encabezado)
    if not n:
        return None
    if "tachado" in n or "original" in n or "lista" in n:
        return "tachado"
    if "final" in n or "oferta" in n or n in ("precio", "deal price", "precio promo"):
        return "final"
    if n in ("mla", "item", "item id", "id", "publicacion") or n.startswith("mla"):
        return "mla"
    if n == "sku" or n.startswith("sku"):
        return "sku"
    return None


def _precio(v) -> Decimal | None:
    if v is None or str(v).strip() == "":
        return None
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v))
    t = re.sub(r"[^\d,.\-]", "", str(v))
    if not t:
        return None
    if "," in t and "." in t:                       # 1.234,56 (AR) o 1,234.56 (US): el último separador es el decimal
        t = t.replace(".", "").replace(",", ".") if t.rfind(",") > t.rfind(".") else t.replace(",", "")
    elif "," in t:
        t = t.replace(",", ".")
    elif t.count(".") > 1 or re.fullmatch(r"-?\d{1,3}\.\d{3}", t):  # 12.500 → miles
        t = t.replace(".", "")
    try:
        return Decimal(t)
    except InvalidOperation:
        return None


def normalizar_mla(v) -> str | None:
    t = str(v or "").strip()
    m = _MLA.search(t)
    if m:
        return "MLA" + m.group(1)
    return None


def leer_excel(contenido: bytes) -> dict:
    """Devuelve `{"filas": [...], "errores": [...]}`. Cada fila: fila (n.º en el Excel), sku, mla, tachado, final y
    `bloqueos` (lo que ya se ve mal solo con el Excel). Errores de formato (columnas que faltan, archivo ilegible) van
    en `errores` y no hay filas."""
    from openpyxl import load_workbook

    try:
        ws = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True).worksheets[0]
    except Exception as e:  # noqa: BLE001 — cualquier archivo que no sea un xlsx legible
        return {"filas": [], "errores": [f"No se pudo leer el archivo como Excel (.xlsx): {e}"]}
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"filas": [], "errores": ["El Excel está vacío."]}
    cols: dict[str, int] = {}
    for i, h in enumerate(rows[0]):
        c = _columna(h)
        if c and c not in cols:
            cols[c] = i
    faltan = [n for n, k in (("SKU", "sku"), ("MLA", "mla"), ("precio tachado", "tachado"), ("precio final", "final")) if k not in cols]
    if faltan:
        return {"filas": [], "errores": [f"Faltan columnas en la primera fila: {', '.join(faltan)}. "
                                          "Tienen que ser: SKU, MLA, precio tachado, precio final."]}
    filas, vistos = [], {}
    for n, r in enumerate(rows[1:], start=2):
        def cel(k):
            return r[cols[k]] if cols[k] < len(r) else None
        if all(x is None or str(x).strip() == "" for x in (cel("sku"), cel("mla"), cel("tachado"), cel("final"))):
            continue
        mla, tachado, final = normalizar_mla(cel("mla")), _precio(cel("tachado")), _precio(cel("final"))
        bloqueos: list[str] = []
        if not mla:
            bloqueos.append(f"El MLA «{cel('mla') or ''}» no es válido (tiene que ser MLA seguido de números).")
        if tachado is None or tachado <= 0:
            bloqueos.append("Falta el precio tachado o no es un número mayor a 0.")
        if final is None or final <= 0:
            bloqueos.append("Falta el precio final o no es un número mayor a 0.")
        if tachado and final and tachado > 0 and final > 0:
            if final >= tachado:
                bloqueos.append(f"El precio final (${final:,.0f}) tiene que ser menor al tachado (${tachado:,.0f}): así no hay oferta.".replace(",", "."))
            else:
                desc = (1 - final / tachado) * 100
                if desc < DESCUENTO_MIN_PCT or desc > DESCUENTO_MAX_PCT:
                    bloqueos.append(f"El descuento es {desc:.1f} %: las campañas propias solo admiten de {DESCUENTO_MIN_PCT} % a {DESCUENTO_MAX_PCT} %.")
        if mla:
            if mla in vistos:
                bloqueos.append(f"El MLA está repetido en el Excel (también en la fila {vistos[mla]}).")
            else:
                vistos[mla] = n
        filas.append({"fila": n, "sku": str(cel("sku") or "").strip(), "mla": mla or str(cel("mla") or "").strip(),
                      "tachado": float(tachado) if tachado else None, "final": float(final) if final else None,
                      "bloqueos": bloqueos, "avisos": []})
    if not filas:
        return {"filas": [], "errores": ["No hay filas con datos debajo de los encabezados."]}
    if len(filas) > MAX_FILAS:
        return {"filas": [], "errores": [f"El Excel tiene {len(filas)} filas; el máximo por carga es {MAX_FILAS}."]}
    return {"filas": filas, "errores": []}


# ── Análisis (solo lectura) ──

def _entrada_promo(p: dict) -> tuple[str | None, str | None]:
    return p.get("promotion_id") or p.get("id"), p.get("status")


def analizar(
    ml: MLOfertasClient, cuenta: str, filas: list[dict], costo_provider=None, iva_provider=None, tc: Decimal | None = None,
    params: ParametrosMargen | None = None, progreso_cb: Callable[[int, int, str], None] | None = None,
) -> dict:
    """Solo lectura. Completa `bloqueos`/`avisos` de cada fila, agrega `candidatas` (ids de campañas donde el MLA puede
    entrar) y arma `campanas`: las campañas propias de la cuenta con cuántas filas de las que se pueden pueden participar."""
    params = params or ParametrosMargen()
    validas = [f for f in filas if f["mla"] and not any("no es válido" in b for b in f["bloqueos"])]
    ids = list(dict.fromkeys(f["mla"] for f in validas))
    if progreso_cb:
        progreso_cb(0, len(ids) * 2, "Leyendo las publicaciones")
    detalles = {d.get("id"): d for d in ml.detalle_items_ofertas(ids, cuenta)} if ids else {}

    def _promos(item_id: str):
        try:
            return item_id, ml.promociones_item(item_id, cuenta), None
        except Exception as e:  # noqa: BLE001
            return item_id, [], str(e)

    promos: dict[str, tuple[list, str | None]] = {}
    ids_existentes = [i for i in ids if i in detalles]
    with ThreadPoolExecutor(max_workers=5) as ex:
        for k, (item_id, lista, err) in enumerate(ex.map(_promos, ids_existentes), start=1):
            promos[item_id] = (lista or [], err)
            if progreso_cb:
                progreso_cb(len(ids) + k, len(ids) * 2, "Mirando en qué campañas puede entrar cada una")

    campanas_cuenta = [c for c in ml.promociones_seller(cuenta) if c.get("type") == "SELLER_CAMPAIGN" and c.get("id")]
    for f in filas:
        f.setdefault("candidatas", [])
        f.setdefault("avisos", [])
        d = detalles.get(f["mla"])
        if not f["mla"] or any("no es válido" in b for b in f["bloqueos"]):
            continue
        if d is None:
            f["bloqueos"].append(f"{f['mla']} no existe en la cuenta {cuenta} (o es de la otra cuenta): revisá que el Excel sea de esta cuenta.")
            continue
        f["titulo"] = d.get("title", "")
        f["precio_actual"] = d.get("price")
        sku_ml = _sku_de_item(d)
        f["sku_ml"] = sku_ml
        if sku_ml and f["sku"] and _norm(sku_ml).replace(" ", "") != _norm(f["sku"]).replace(" ", ""):
            f["avisos"].append(f"El SKU del Excel ({f['sku']}) no coincide con el de la publicación en ML ({sku_ml}): ¿es el MLA correcto?")
        if not sku_ml:
            f["avisos"].append("La publicación no tiene SKU cargado en ML: no se pudo cotejar con el del Excel.")
        if d.get("status") and d["status"] != "active":
            f["avisos"].append(f"La publicación está «{d['status']}» en ML ({d.get('sub_status') or 'sin detalle'}).")
        if d.get("available_quantity") == 0:
            f["avisos"].append("La publicación no tiene stock.")
        lista, err = promos.get(f["mla"], ([], None))
        if err:
            f["avisos"].append(f"No se pudieron leer sus promociones en ML ({err}): no sé en qué campañas puede entrar.")
        for p in lista:
            pid, st = _entrada_promo(p)
            if pid and p.get("type") == "SELLER_CAMPAIGN" and st in ("candidate", "started", "pending"):
                f["candidatas"].append({"id": pid, "status": st, "min": p.get("min_discounted_price"), "max": p.get("max_discounted_price")})
        # margen: solo avisa
        f["margen_pct"] = None
        if costo_provider and iva_provider and tc and sku_ml and f["final"]:
            costo_usd, iva = costo_provider.obtener(sku_ml), iva_provider.factor(sku_ml)
            if costo_usd is None or iva is None:
                f["avisos"].append("No pude calcular el margen: falta el costo o el IVA del SKU.")
            else:
                r = calcular_margen_oferta(Decimal(str(f["final"])), iva, costo_usd * tc, d.get("domain_id"), _cuotas_sin_interes(d), params)
                f["margen_pct"] = float(r.margen_pct * 100) if r.margen_pct is not None else None
                f["margen_ars"] = float(r.margen)
                if r.margen < 0:
                    f["avisos"].append(f"Con ese precio final se pierde plata: margen {f['margen_pct']:.1f} % (${float(r.margen):,.0f} por venta).".replace(",", "."))

    ok = [f for f in filas if not f["bloqueos"]]
    campanas = []
    for c in campanas_cuenta:
        cid = c["id"]
        pueden = [f["mla"] for f in ok if any(x["id"] == cid for x in f["candidatas"])]
        campanas.append({
            "id": cid, "nombre": c.get("name") or cid, "status": c.get("status"),
            "fecha_desde": c.get("date_from") or c.get("start_date"), "fecha_hasta": c.get("date_to") or c.get("finish_date") or c.get("end_date"),
            "pueden": len(pueden), "mlas": pueden,
        })
    campanas.sort(key=lambda c: (-c["pueden"], c["nombre"]))
    return {"filas": filas, "campanas": campanas, "total": len(filas), "sin_bloqueo": len(ok)}


def iniciar_analisis(job_id: str, cuenta: str, filas: list[dict], costo_provider=None, iva_provider=None, tc: Decimal | None = None,
                     ml: MLOfertasClient | None = None) -> None:
    _jobs[job_id] = {"status": "running", "progress": None, "result": None, "log": []}

    def _prog(a, t, label):
        _jobs[job_id]["progress"] = {"current": a, "total": t, "label": label}

    try:
        _jobs[job_id]["result"] = analizar(ml or MLOfertasClient(), cuenta, filas, costo_provider, iva_provider, tc, progreso_cb=_prog)
        _jobs[job_id]["status"] = "done"
    except Exception as e:  # noqa: BLE001
        _jobs[job_id]["status"] = "error"
        _jobs[job_id]["log"].append(f"Error: {e}")
    _jobs[job_id]["progress"] = None


# ── Aplicación (ESCRIBE en Mercado Libre) ──

def _chequear_fila(f: dict, promotion_id: str) -> str | None:
    """Motivo por el que la fila NO debe intentarse en esta campaña (sin llamar a ML), o None."""
    if f.get("bloqueos"):
        return "No se puede: " + " ".join(f["bloqueos"])
    if not (f.get("final") and f.get("tachado") and 0 < f["final"] < f["tachado"]):
        return "No se puede: el precio final tiene que ser mayor a 0 y menor al tachado."
    cand = next((x for x in f.get("candidatas", []) if x["id"] == promotion_id), None)
    if not cand:
        return "Esta publicación no puede participar en esa campaña."
    if cand.get("min") is not None and f["final"] < float(cand["min"]):
        return f"No se puede ir a ese precio: el mínimo que acepta ML en esta campaña es ${float(cand['min']):,.0f}.".replace(",", ".")
    return None


def aplicar(
    ml: MLOfertasEscritura, cuenta: str, promotion_id: str, filas: list[dict], operador: str = "",
    progreso_cb: Callable[[int, int, str], None] | None = None, historial_cb: Callable[[list[dict]], None] | None = None,
) -> dict:
    """Una por una, con el tachado y el final del Excel. Si ML pide una verificación de seguridad (QR/cara) se corta todo
    el lote: ninguna de las que siguen puede pasar sin vos. Devuelve el resultado por fila y las que quedan pendientes."""
    resultados: list[dict] = []
    historial: list[dict] = []
    cortado = False
    total = len(filas)
    for i, f in enumerate(filas, start=1):
        if progreso_cb:
            progreso_cb(i - 1, total, f"Aplicando {f['mla']}")
        base = {"fila": f["fila"], "mla": f["mla"], "sku": f.get("sku", "")}
        if cortado:
            resultados.append({**base, "estado": "pendiente", "mensaje": "No se intentó: ML pidió una verificación de seguridad antes."})
            continue
        motivo = _chequear_fila(f, promotion_id)
        if motivo:
            resultados.append({**base, "estado": "pendiente", "mensaje": motivo})
            continue
        try:
            r = activar_en_campana_tradicional(
                ml, cuenta, f["mla"], Decimal(str(f["final"])), promotion_id=promotion_id, tachado_exacto=Decimal(str(f["tachado"])))
        except Exception as e:  # noqa: BLE001
            r = {"ok": False, "error": f"Error inesperado: {e}"}
        if r.get("ok"):
            resultados.append({**base, "estado": "ok", "mensaje": "Aplicada.", "precio": r.get("precio"), "tachado": r.get("original_price")})
            historial.append({
                "Fecha": _hoy(), "Cuenta": cuenta, "Item ID": f["mla"], "SKU": f.get("sku", ""), "Accion": "Activar en campaña",
                "Tipo/Promocion": f"SELLER_CAMPAIGN ({r.get('nombre_campana') or promotion_id})", "Precio Anterior": f.get("precio_actual", ""),
                "Precio Nuevo": f"{r.get('precio')} (tachado {r.get('original_price')})",
                "Margen % Resultante": round(f["margen_pct"], 1) if f.get("margen_pct") is not None else "",
                "Bajo Umbral": "Sí" if (f.get("margen_pct") is not None and f["margen_pct"] < 0) else "No",
                "Operador": operador, "Origen": "Excel (lote)",
            })
        else:
            resultados.append({**base, "estado": "error", "mensaje": r.get("error") or "ML no la aceptó."})
            if r.get("requiere_verificacion_manual"):
                cortado = True
    if progreso_cb:
        progreso_cb(total, total, "Listo")
    if historial and historial_cb:
        try:
            historial_cb(historial)
        except Exception as e:  # noqa: BLE001 — el historial no debe tapar el resultado real
            for r in resultados:
                if r["estado"] == "ok":
                    r["mensaje"] += " (No se pudo anotar en Historial Ofertas ML.)"
            return {"resultados": resultados, "aviso_historial": str(e), "cortado_por_verificacion": cortado}
    return {"resultados": resultados, "cortado_por_verificacion": cortado}


def _hoy() -> str:
    from datetime import date
    return date.today().isoformat()


def iniciar_aplicacion(job_id: str, cuenta: str, promotion_id: str, filas: list[dict], operador: str = "",
                       historial_cb: Callable[[list[dict]], None] | None = None, ml: MLOfertasEscritura | None = None) -> None:
    _jobs[job_id] = {"status": "running", "progress": None, "result": None, "log": []}

    def _prog(a, t, label):
        _jobs[job_id]["progress"] = {"current": a, "total": t, "label": label}

    try:
        _jobs[job_id]["result"] = aplicar(ml or MLOfertasEscritura(), cuenta, promotion_id, filas, operador, _prog, historial_cb)
        _jobs[job_id]["status"] = "done"
    except Exception as e:  # noqa: BLE001
        _jobs[job_id]["status"] = "error"
        _jobs[job_id]["log"].append(f"Error: {e}")
    _jobs[job_id]["progress"] = None


def lanzar(fn: Callable, *args, **kw) -> None:
    """Corre un job en un hilo propio (no bloquea el event loop de uvicorn)."""
    threading.Thread(target=fn, args=args, kwargs=kw, daemon=True).start()
